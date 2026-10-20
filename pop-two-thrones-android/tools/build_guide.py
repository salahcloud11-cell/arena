#!/usr/bin/env python3
"""
Build the player guide for "Prince of Persia: The Two Thrones" on Android
(Winlator) in two formats from ONE source of content:

  docs/PoP-Two-Thrones-Android-Guide.docx   (uploaded to Google Drive as a Google Doc)
  README.md                                 (same text, for GitHub)

Simple English on purpose: short sentences, numbered steps, exact Winlator
11.2 menu names (taken from the Winlator source code).

Usage:  python3 build_guide.py      (needs: pip install python-docx)
"""
from __future__ import annotations

import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.normpath(os.path.join(HERE, ".."))
IMG = os.path.join(ROOT, "docs", "touch-controls-layout.png")
DOCX_OUT = os.path.join(ROOT, "docs", "PoP-Two-Thrones-Android-Guide.docx")
MD_OUT = os.path.join(ROOT, "README.md")

BAT = "00 ANDROID - PLAY GAME.bat"
ICP = "00 ANDROID - touch controls.icp"
GUIDE = "00 ANDROID - HOW TO PLAY (guide)"
PROFILE = "Prince of Persia - The Two Thrones"
FOLDER = "princ of percia"
WINLATOR_URL = "https://github.com/brunodev85/winlator/releases"

# --------------------------------------------------------------------------
# Content.  **bold** is supported inside text.  URLs become links.
# Block kinds: title, subtitle, h2, h3, p, steps, bullets, table, callout, image
# --------------------------------------------------------------------------
CONTENT = [
    ("title", "Prince of Persia: The Two Thrones on Android"),
    ("subtitle", "Simple step-by-step guide - Winlator 11.2 - October 2026"),

    ("callout", "important", "Can this game become an Android APK?",
     ["**No.** This is a Windows PC game from 2005. It is made only of Windows program files "
      "(.exe and .dll), and there is no source code. So nobody can convert it into a real Android APK.",
      "**But you CAN play it on your phone.** Winlator is a free Android app (an APK) that runs "
      "Windows games. You install Winlator, copy the game folder to your phone, and start the game "
      "inside Winlator. This guide shows every step."]),

    ("p", "**Files I added to the game folder for you:**"),
    ("bullets", [
        f"**{BAT}** - starts the game with one tap. It also sets the registry keys that the game needs.",
        f"**{ICP}** - on-screen touch buttons made for this game.",
        f"**{GUIDE}** - this guide.",
    ]),

    ("h2", "What you need"),
    ("bullets", [
        "An Android phone with a 64-bit processor. **Snapdragon** phones (Adreno graphics) work best. "
        "Other phones (MediaTek, Exynos, Google Tensor) can also work.",
        "**4 GB RAM** or more (6 - 8 GB is better).",
        "About **5 GB of free space**. The game is about 1.5 GB.",
        "Optional: a **Bluetooth game controller**. It makes the game much easier to play.",
    ]),

    ("h2", "Step 1 - Install Winlator"),
    ("steps", [
        f"On your phone, open: {WINLATOR_URL}",
        "Under **Assets**, tap **Winlator_11.2.apk** (about 157 MB) to download it. "
        "If you see a newer version, take the newest one.",
        "Open the downloaded file and tap **Install**. If Android asks, allow **Install unknown apps**.",
        "Open Winlator and wait until **Installing System Files...** is finished.",
    ]),
    ("callout", "warning", "Be careful",
     ["Download Winlator only from the GitHub page above. Many \"Prince of Persia APK\" files on "
      "other websites are fake or contain viruses."]),

    ("h2", "Step 2 - Copy the game to your phone"),
    ("p", f"The game folder is **{FOLDER}** in your Google Drive (about 1.5 GB). "
          "Copy the whole folder to your phone."),
    ("h3", "Way A - with a computer (easiest)"),
    ("steps", [
        "On the computer, open https://drive.google.com",
        f"Right-click the folder **{FOLDER}** and choose **Download**. Google makes a .zip file "
        "(a big folder can become 2 zip files).",
        "Connect your phone with a USB cable. Copy the zip file(s) into the phone's **Download** folder.",
    ]),
    ("h3", "Way B - phone only"),
    ("steps", [
        "Open **Chrome** and go to drive.google.com. Tap **\u22ee** (the three dots, top right) and turn on **Desktop site**.",
        f"Long-press the folder **{FOLDER}** and choose **Download**.",
    ]),
    ("h3", "Then unzip it"),
    ("steps", [
        "Open the zip file with a file app, for example **Files by Google** or **ZArchiver**, and tap "
        "**Extract**. If there are 2 zip files, extract both into the same place.",
        f"Check that the game files are here: **Internal storage > Download > {FOLDER} > "
        "PrinceOfPersia.exe, pop3.exe, pop3.bf ...**",
        f"If you see a **{FOLDER}** folder inside another **{FOLDER}** folder, move the inner one up "
        "into **Download**.",
    ]),
    ("callout", "tip", "Good to know",
     [f"Winlator shows your **Download** folder as drive **D:**. So the game will be at "
      f"**D:\\{FOLDER}**."]),

    ("h2", "Step 3 - Create a container (a small \"Windows PC\" inside Winlator)"),
    ("steps", [
        "Open Winlator. Tap **\u2630** (menu) > **Containers** > **+** (add).",
        "Set these options:",
    ]),
    ("table", ["Setting", "Choose this"], [
        ["Name", "Prince of Persia"],
        ["Screen Size", "**1280x720 (16:9)**. Older or slower phone: **854x480 (16:9)**"],
        ["Wine Version", "Leave it as it is"],
        ["Graphics Driver", "**Vulkan: Turnip** - Snapdragon phones.\n**Vulkan: Vortek** - all other phones.\nLeave **OpenGL** as it is."],
        ["DX Wrapper", "**Direct3D: DXVK**. Then tap the settings button on the **Direct3D** row:\n"
                       "**Frame Rate = 60** (very important). Keep **Version** as it is. Tap **OK**.\n"
                       "Leave **DirectX 12** as it is."],
        ["Audio Driver", "**ALSA** (if you have no sound later: **PulseAudio**)"],
        ["HUD Mode", "**Simple FPS** (optional - shows the game speed)"],
        ["Advanced tab > Box64 Preset", "**Performance** (if the game crashes: **Intermediate**, then **Stability**)"],
    ]),
    ("steps_continue", [
        "Tap **\u2713** to save. Wait until the container is ready.",
    ], 3),
    ("callout", "important", "Why Frame Rate 60?",
     ["This game has bugs above 60 FPS: it can freeze on loading screens, jumps can fail, and the "
      "sound can go out of sync. Frame Rate 60 stops this."]),

    ("h2", "Step 4 - Add the touch controls"),
    ("steps", [
        "In Winlator, tap **\u2630** > **Input Controls**.",
        "Tap the **Import** button > **Open File**.",
        f"Go to **Download > {FOLDER}** and choose **{ICP}**.",
        f"The profile **{PROFILE}** is now in the list.",
    ]),

    ("h2", "Step 5 - Start the game"),
    ("steps", [
        "Tap **\u2630** > **Containers**. Next to **Prince of Persia**, tap **\u22ee** (three dots) > **File Manager**.",
        f"Open **D:**, then open **{FOLDER}** (tap the folder icon or its **Open** button).",
        f"Find **{BAT}** and tap its **Run** button (or its icon). A black window shows for a moment, "
        "then the game launcher opens.",
        "Tap **Play** in the launcher. The first time, the launcher checks the system - just wait a few seconds.",
    ]),
    ("callout", "tip", "Faster next time",
     [f"Tap **\u22ee** (three dots) next to **{BAT}** > **Add to Favorites**. Next time open "
      "**File Manager > Favorites** and tap it."]),

    ("h2", "Step 6 - Show the buttons on the screen"),
    ("steps", [
        "While the game is running, press your phone's **Back** button (or swipe in from the left edge). "
        "The Winlator menu opens.",
        f"Tap **Input Controls**, choose **{PROFILE}**, turn on **Show Touchscreen Controls**, and tap **OK**.",
        "The first time, the game asks for a **player name**. Open the Winlator menu again > **Keyboard**, "
        "type a name, and press **Enter**.",
    ]),

    ("h2", "Step 7 - Settings inside the game (Options)"),
    ("bullets", [
        "**Graphics - Basic > Resolution:** the same as your container (for example 1280x720), or the closest one.",
        "**Graphics - Advanced > V-Sync:** On.",
        "If the game is slow: **Shadow Effects** low or off, **Water Effects** off.",
    ]),

    ("h2", "The touch buttons"),
    ("image", IMG, "Yellow text = notes for you. It is not shown in the game."),
    ("table", ["Button", "What it does", "PC key"], [
        ["MOVE (left stick)", "Walk and run. Also moves in menus.", "W A S D"],
        ["CAMERA (right stick)", "Turns the camera.", "Mouse"],
        ["JUMP", "Jump, roll, climb up.", "Space"],
        ["ATTACK", "Dagger attack, finish a Speed Kill.", "Left mouse"],
        ["E-KILL", "Start a Speed Kill, use the second weapon, grab.", "E"],
        ["C-PICK", "Pick up or throw a weapon, climb down.", "C"],
        ["WALL", "Run on walls, swing on bars, push / pull, drink water, block.", "Right mouse"],
        ["REWIND", "Hold = rewind time. Tap = slow motion.", "R"],
        ["F-LOOK", "Look around (first-person view).", "F"],
        ["Q-VIEW", "Wide view of the area.", "Q"],
        ["X-CAM", "Put the camera behind the Prince.", "X"],
        ["ESC", "Pause menu / go back.", "Esc"],
        ["ENTER", "OK / select in menus.", "Enter"],
    ]),
    ("bullets", [
        "The game's help messages show **PC keys** (for example \"Space\" or \"E\"). Use the table to find the button.",
        "In menus: choose with the **MOVE** stick, press **ENTER** to select, **ESC** to go back.",
        "You can also swipe on an empty part of the screen to turn the camera. "
        "(A short tap on an empty part = attack.)",
        "To move or resize buttons: **\u2630 > Input Controls** > choose the profile > **Controls Editor**.",
        "Camera too slow or too fast? **\u2630 > Input Controls** > change **Cursor Speed**.",
    ]),

    ("h2", "Playing with a game controller (optional)"),
    ("steps", [
        "Pair the controller in Android **Settings > Bluetooth** first.",
        "Start the game. In the game, open **Options > Controls - Gamepad** and set the buttons you like.",
        "If the game does not react to the controller: open the container settings > **Wine Configuration** > "
        "**DirectInput Mapper Type** and try the other choice.",
        "On Xbox-style controllers this game cannot use the two triggers (LT / RT) correctly. "
        "Put important actions on other buttons.",
    ]),

    ("h2", "Problems and fixes"),
    ("p", "\"DXVK settings\" = container settings > **DX Wrapper** > the settings button on the **Direct3D** row."),
    ("table", ["Problem", "Try this"], [
        ["Black screen, or the game closes at the start",
         "DXVK settings > **Version 1.10.3** (or 1.7.2).\nOr switch the **Vulkan** driver (Turnip <> Vortek).\n"
         "Or **Box64 Preset: Intermediate**, then **Stability**."],
        ["The launcher says your video card is not supported",
         "DXVK settings > **Custom Device** > choose an NVIDIA card, for example **NVIDIA GeForce GTX 480**."],
        ["Freezes on a loading screen, jumps fail, sound is out of sync",
         "Check DXVK settings > **Frame Rate = 60**. In the game, turn **V-Sync** on."],
        ["The game is slow",
         "**Screen Size 854x480** or **800x600** (and the same resolution in the game). "
         "Turn off shadows and water effects. Close other apps."],
        ["No sound, or crackling sound", "**Audio Driver: PulseAudio**."],
        ["The camera turns by itself, or does not turn",
         "Container settings > **Wine Configuration > Mouse Warp Override > Force** (or Enable).\n"
         "Or, while playing: Winlator menu > **Input Controls** > turn on **Relative Mouse Movement**."],
        ["Nothing happens when you tap the .bat file",
         "Tap **RUN THIS FIX FIRST.bat** once. Wait 30 seconds, then close it. "
         "Now tap **PrinceOfPersia.exe** to play."],
        ["\"PrinceOfPersia.exe was not found\"",
         "The .bat file must stay inside the game folder, next to PrinceOfPersia.exe."],
    ]),

    ("h2", "Good to know"),
    ("bullets", [
        f"Your **saved games** are kept inside the game folder (Download > {FOLDER}). Do not delete this folder.",
        "Players report: **Snapdragon 8 Elite** - stable 60 FPS. **Snapdragon 870** - playable. "
        "Even a mid-range **MediaTek Helio G99** phone (8 GB RAM) was playable at 800x600 in a test.",
        "You can ignore these files on the phone: \"Launch Prince of Persia The Two Thrones.lnk\", "
        "\"AnkerGames - Free Pre-installed PC Games.url\", \"Read Me.txt\".",
        "More help about Winlator: https://github.com/brunodev85/winlator",
    ]),
]

URL_RE = re.compile(r"(https?://[^\s)]+)")


# ----------------------------------------------------------------- Markdown
def md_inline(text: str) -> str:
    text = text.replace("\n", "<br>")
    return URL_RE.sub(lambda m: f"<{m.group(1)}>", text)


def build_markdown() -> str:
    out = []
    step_no = 0
    for block in CONTENT:
        kind = block[0]
        if kind == "title":
            out.append(f"# {block[1]}\n")
        elif kind == "subtitle":
            out.append(f"_{block[1]}_\n")
        elif kind == "h2":
            out.append(f"## {block[1]}\n")
        elif kind == "h3":
            out.append(f"**{block[1]}**\n")
        elif kind == "p":
            out.append(md_inline(block[1]) + "\n")
        elif kind in ("steps", "steps_continue"):
            start = block[2] if kind == "steps_continue" else 1
            for i, item in enumerate(block[1], start):
                out.append(f"{i}. {md_inline(item)}")
            out.append("")
        elif kind == "bullets":
            for item in block[1]:
                out.append(f"- {md_inline(item)}")
            out.append("")
        elif kind == "table":
            header, rows = block[1], block[2]
            out.append("| " + " | ".join(header) + " |")
            out.append("|" + "|".join("---" for _ in header) + "|")
            for r in rows:
                out.append("| " + " | ".join(md_inline(c) for c in r) + " |")
            out.append("")
        elif kind == "callout":
            icon = {"important": "\u2757", "warning": "\u26a0\ufe0f", "tip": "\U0001f4a1"}[block[1]]
            out.append(f"> {icon} **{block[2]}**")
            for para in block[3]:
                out.append(">")
                out.append(f"> {md_inline(para)}")
            out.append("")
        elif kind == "image":
            rel = os.path.relpath(block[1], ROOT).replace(os.sep, "/")
            out.append(f"![Touch controls layout]({rel})\n")
            out.append(f"_{block[2]}_\n")
    out.append("---\n")
    out.append("### Files in this folder (for developers)\n")
    out.append("- `winlator/" + ICP + "` - Winlator input-controls profile (import in **Input Controls > Import > Open File**).")
    out.append("- `winlator/" + BAT + "` - one-tap start script: writes the launcher's registry keys, then starts `PrinceOfPersia.exe`.")
    out.append("- `tools/make_controls.py` - generates the `.icp` (validated against Winlator's binding names, overlap-checked for 16:9 - 21:9 screens) and the preview image.")
    out.append("- `tools/build_guide.py` - builds this README and the Word/Google Docs guide from one source.")
    out.append("- `docs/` - generated guide (`.docx`) and layout preview.\n")
    out.append("The game files themselves are **not** in this repository.\n")
    return "\n".join(out)


# --------------------------------------------------------------------- DOCX
def build_docx(path: str) -> None:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    ACCENT = RGBColor(0x8A, 0x4B, 0x16)   # desert brown
    DARK = RGBColor(0x22, 0x22, 0x22)
    CALLOUT = {
        "important": ("FFF4CC", "E0A800"),
        "warning": ("FDE2E1", "D93025"),
        "tip": ("E6F4EA", "1E8E3E"),
    }

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(1.8))
    sec.top_margin = sec.bottom_margin = Cm(1.6)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(11.5)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.15

    def shade(cell, fill_hex):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill_hex)
        tcPr.append(shd)

    def cell_borders(cell, color_hex, left_only=False):
        tcPr = cell._tc.get_or_add_tcPr()
        borders = OxmlElement("w:tcBorders")
        for edge in ("top", "left", "bottom", "right"):
            el = OxmlElement(f"w:{edge}")
            if left_only and edge != "left":
                el.set(qn("w:val"), "nil")
            else:
                el.set(qn("w:val"), "single")
                el.set(qn("w:sz"), "24" if edge == "left" else "4")
                el.set(qn("w:color"), color_hex)
            borders.append(el)
        tcPr.append(borders)

    def add_hyperlink(paragraph, url, text, bold=False):
        part = paragraph.part
        r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                              is_external=True)
        hl = OxmlElement("w:hyperlink")
        hl.set(qn("r:id"), r_id)
        run = OxmlElement("w:r")
        rPr = OxmlElement("w:rPr")
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "1A73E8")
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        rPr.append(color)
        rPr.append(u)
        if bold:
            rPr.append(OxmlElement("w:b"))
        run.append(rPr)
        t = OxmlElement("w:t")
        t.text = text
        t.set(qn("xml:space"), "preserve")
        run.append(t)
        hl.append(run)
        paragraph._p.append(hl)

    def add_rich(paragraph, text, base_bold=False, color=None, size=None):
        """Add text with **bold** markup, URLs as links and \\n as line breaks."""
        parts = re.split(r"(\*\*.+?\*\*)", text)
        for part in parts:
            if not part:
                continue
            bold = base_bold
            if part.startswith("**") and part.endswith("**"):
                part, bold = part[2:-2], True
            lines = part.split("\n")
            for li, line in enumerate(lines):
                if li > 0:
                    paragraph.add_run().add_break()
                for chunk in URL_RE.split(line):
                    if not chunk:
                        continue
                    if URL_RE.fullmatch(chunk):
                        add_hyperlink(paragraph, chunk, chunk, bold)
                    else:
                        run = paragraph.add_run(chunk)
                        run.bold = bold
                        if color is not None:
                            run.font.color.rgb = color
                        if size is not None:
                            run.font.size = size

    def numbered(items, start=1):
        for i, item in enumerate(items, start):
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.left_indent = Cm(0.9)
            pf.first_line_indent = Cm(-0.7)
            pf.space_after = Pt(5)
            r = p.add_run(f"{i}.  ")
            r.bold = True
            r.font.color.rgb = ACCENT
            add_rich(p, item)

    def bulleted(items):
        for item in items:
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.left_indent = Cm(0.9)
            pf.first_line_indent = Cm(-0.5)
            pf.space_after = Pt(4)
            r = p.add_run("\u2022  ")
            r.bold = True
            r.font.color.rgb = ACCENT
            add_rich(p, item)

    for block in CONTENT:
        kind = block[0]
        if kind == "title":
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(block[1])
            r.bold = True
            r.font.size = Pt(22)
            r.font.color.rgb = ACCENT
        elif kind == "subtitle":
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(10)
            r = p.add_run(block[1])
            r.italic = True
            r.font.size = Pt(11)
            r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        elif kind == "h2":
            h = doc.add_heading(level=2)
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(6)
            r = h.add_run(block[1])
            r.font.size = Pt(15.5)
            r.font.color.rgb = ACCENT
            r.font.name = "Arial"
        elif kind == "h3":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            r = p.add_run(block[1])
            r.bold = True
            r.font.size = Pt(12)
            r.font.color.rgb = DARK
        elif kind == "p":
            p = doc.add_paragraph()
            add_rich(p, block[1])
        elif kind == "steps":
            numbered(block[1])
        elif kind == "steps_continue":
            numbered(block[1], block[2])
        elif kind == "bullets":
            bulleted(block[1])
        elif kind == "table":
            header, rows = block[1], block[2]
            t = doc.add_table(rows=1, cols=len(header))
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for i, h in enumerate(header):
                c = t.rows[0].cells[i]
                c.text = ""
                add_rich(c.paragraphs[0], h, base_bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
                shade(c, "8A4B16")
            for ri, row in enumerate(rows):
                cells = t.add_row().cells
                for i, val in enumerate(row):
                    cells[i].text = ""
                    add_rich(cells[i].paragraphs[0], val, base_bold=(i == 0 and len(header) == 3))
                    if ri % 2 == 1:
                        shade(cells[i], "F7F1EA")
            if len(header) == 3:
                widths = [Cm(4.2), Cm(10.0), Cm(3.2)]
            else:
                widths = [Cm(5.6), Cm(11.8)]
            # Google Docs uses the table grid (gridCol) widths, Word uses cell widths: set both.
            t.autofit = False
            for i, w in enumerate(widths):
                t.columns[i].width = w
            for row in t.rows:
                for i, w in enumerate(widths):
                    row.cells[i].width = w
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
        elif kind == "callout":
            fill, border = CALLOUT[block[1]]
            t = doc.add_table(rows=1, cols=1)
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            t.autofit = False
            t.columns[0].width = Cm(17.4)
            c = t.rows[0].cells[0]
            c.width = Cm(17.4)
            shade(c, fill)
            cell_borders(c, border)
            p = c.paragraphs[0]
            icon = {"important": "\u2757 ", "warning": "\u26a0 ", "tip": "\U0001f4a1 "}[block[1]]
            r = p.add_run(icon + block[2])
            r.bold = True
            r.font.size = Pt(12)
            for para in block[3]:
                pp = c.add_paragraph()
                pp.paragraph_format.space_after = Pt(3)
                add_rich(pp, para)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
        elif kind == "image":
            if os.path.exists(block[1]):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(block[1], width=Cm(17.0))
                cap = doc.add_paragraph()
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = cap.add_run(block[2])
                r.italic = True
                r.font.size = Pt(10)
                r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    doc.core_properties.title = "Prince of Persia: The Two Thrones on Android (Winlator guide)"
    doc.core_properties.author = "Arena agent"
    doc.save(path)


def main() -> int:
    md = build_markdown()
    with open(MD_OUT, "w", encoding="utf-8") as fh:
        fh.write(md)
    print("wrote", MD_OUT)
    try:
        build_docx(DOCX_OUT)
        print("wrote", DOCX_OUT)
    except ImportError:
        print("python-docx not installed - skipped .docx")
    return 0


if __name__ == "__main__":
    sys.exit(main())
